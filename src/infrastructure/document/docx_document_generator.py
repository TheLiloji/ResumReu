"""DocumentGeneratorPort implementation producing a .docx file via python-docx."""

from __future__ import annotations

import logging
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from src.application.ports.document_generator_port import DocumentGeneratorPort
from src.domain.entities.document import MeetingDocument

logger = logging.getLogger(__name__)


class DocxDocumentGenerator(DocumentGeneratorPort):
    """Renders a structured MeetingDocument to a Word file."""

    def generate(self, document: MeetingDocument, output_path: Path) -> Path:
        doc = Document()
        self._write_title(doc, document)
        self._write_metadata(doc, document)
        self._write_summary(doc, document)
        self._write_bullets(doc, "Points clés", document.key_points)
        self._write_bullets(doc, "Décisions", document.decisions)
        self._write_action_items(doc, document)
        self._write_technical_terms(doc, document)
        self._write_extra_sections(doc, document)
        doc.save(output_path)
        logger.info("Wrote document %s", output_path)
        return output_path

    @staticmethod
    def _write_title(doc: Document, document: MeetingDocument) -> None:
        title = doc.add_heading(document.title or "Compte rendu de réunion", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    @staticmethod
    def _write_metadata(doc: Document, document: MeetingDocument) -> None:
        meta = doc.add_paragraph()
        meta.add_run("Date : ").bold = True
        meta.add_run(document.date.strftime("%Y-%m-%d %H:%M"))
        participants = doc.add_paragraph()
        participants.add_run("Participants : ").bold = True
        participants.add_run(
            ", ".join(document.participants) if document.participants else "(non détectés)"
        )

    @staticmethod
    def _write_summary(doc: Document, document: MeetingDocument) -> None:
        doc.add_heading("Synthèse", level=1)
        para = doc.add_paragraph(document.summary)
        for run in para.runs:
            run.font.size = Pt(11)

    @staticmethod
    def _write_bullets(doc: Document, title: str, items: tuple[str, ...]) -> None:
        if not items:
            return
        doc.add_heading(title, level=1)
        for item in items:
            doc.add_paragraph(item, style="List Bullet")

    @staticmethod
    def _write_action_items(doc: Document, document: MeetingDocument) -> None:
        if not document.action_items:
            return
        doc.add_heading("Actions", level=1)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Description"
        hdr[1].text = "Responsable"
        hdr[2].text = "Échéance"
        for item in document.action_items:
            row = table.add_row().cells
            row[0].text = item.description
            row[1].text = item.owner or "-"
            row[2].text = item.due_date or "-"

    @staticmethod
    def _write_technical_terms(doc: Document, document: MeetingDocument) -> None:
        if not document.technical_terms:
            return
        doc.add_heading("Glossaire technique", level=1)
        for term, definition in document.technical_terms:
            para = doc.add_paragraph()
            para.add_run(f"{term} : ").bold = True
            para.add_run(definition)

    @staticmethod
    def _write_extra_sections(doc: Document, document: MeetingDocument) -> None:
        for section in document.extra_sections:
            doc.add_heading(section.title, level=1)
            doc.add_paragraph(section.content)
